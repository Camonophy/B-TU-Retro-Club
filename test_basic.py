#!/usr/bin/env python3
"""
Basic test to verify the project structure, imports, and URL generation
work correctly against the current Kleinanzeigen URL scheme.

If any check fails, the script exits with a non-zero status and a
clear summary of which checks failed.
"""

import sys
import os
import traceback

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


_failures = []


def _record(name: str, ok: bool, detail: str = ""):
    status = "✅" if ok else "❌"
    print(f"  {status} {name}" + (f"  ({detail})" if detail else ""))
    if not ok:
        _failures.append(name)


def test_imports():
    """Test that all imports work"""
    print("\n[1] Imports")
    try:
        from config.settings import Settings, settings
        _record("config.settings", True)
    except Exception as e:
        _record("config.settings", False, str(e))
        return

    try:
        from scraper.models import Listing, ScrapeResult
        _record("scraper.models", True)
    except Exception as e:
        _record("scraper.models", False, str(e))

    try:
        from scraper.utils import (
            parse_kleinanzeigen_date,
            calculate_age_days,
            generate_search_url,
            generate_all_category_urls,
            get_random_user_agent,
            build_page_url,
        )
        _record("scraper.utils", True)
    except Exception as e:
        _record("scraper.utils", False, str(e))

    try:
        from scraper.kleinanzeigen import KleinanzeigenScraper, scrape_kleinanzeigen
        _record("scraper.kleinanzeigen", True)
    except Exception as e:
        _record("scraper.kleinanzeigen", False, str(e))

    try:
        from scraper.exporter import ExcelExporter, export_to_excel
        _record("scraper.exporter", True)
    except Exception as e:
        _record("scraper.exporter", False, str(e))


def test_date_parsing():
    """Test date parsing for every format used by Kleinanzeigen."""
    print("\n[2] Date parsing")
    from scraper.utils import parse_kleinanzeigen_date
    from datetime import datetime, timedelta

    cases = [
        # (input, expected_offset_days, description)
        ("Heute",          0,   "today"),
        ("Gestern",        1,   "yesterday"),
        ("vor 2 Tagen",    2,   "N days ago"),
        ("vor 3 Wochen",   21,  "N weeks ago"),
        ("vor 2 Monaten",  None, "N months ago (varies)"),
        ("vor 1 Jahren",   None, "N years ago (varies)"),
        ("01.01.2024",     None, "DD.MM.YYYY"),
        ("Januar 2024",    None, "Month YYYY"),
    ]
    for raw, expected_days, desc in cases:
        parsed = parse_kleinanzeigen_date(raw)
        if expected_days is None:
            # Just check it's parseable
            _record(desc, parsed is not None, f"got {parsed!r}")
        else:
            if parsed is None:
                _record(desc, False, "returned None")
            else:
                delta = (datetime.now() - parsed).days
                _record(desc, delta == expected_days, f"got {delta} days")


def test_settings():
    """Test settings configuration."""
    print("\n[3] Settings")
    from config.settings import Settings
    _record("BASE_URL set", Settings.BASE_URL.startswith("https://"), Settings.BASE_URL)
    _record("REQUEST_DELAY is (min, max) tuple", isinstance(Settings.REQUEST_DELAY, tuple)
            and len(Settings.REQUEST_DELAY) == 2)
    _record("MAX_PAGES > 0", Settings.MAX_PAGES > 0, str(Settings.MAX_PAGES))
    _record("MIN_AGE_DAYS == 90", Settings.MIN_AGE_DAYS == 90, str(Settings.MIN_AGE_DAYS))
    _record("IMMOBILIEN_CATEGORY == c195", Settings.IMMOBILIEN_CATEGORY == "c195",
            Settings.IMMOBILIEN_CATEGORY)
    _record("Output dir exists or can be created", Settings.OUTPUT_DIR.exists())


def test_bundesland_mapping():
    """Test Bundesland mapping has locationIds."""
    print("\n[4] Bundesland mapping")
    import json
    from config.settings import Settings
    try:
        with open(Settings.BUNDESLAND_MAPPING_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except Exception as e:
        _record("load mapping file", False, str(e))
        return

    _record("loaded 16 Bundesländer", len(data) == 16, f"{len(data)} entries")

    missing_loc_id = [name for name, info in data.items() if not info.get("location_id")]
    _record("every entry has a location_id", len(missing_loc_id) == 0,
            "missing: " + ", ".join(missing_loc_id) if missing_loc_id else "")

    # Check that the well-known mapping for NRW, Bayern, Berlin works
    expected = {
        "Bayern": "5510",
        "Berlin": "3331",
        "Nordrhein-Westfalen": "928",
        "Bremen": "1",
    }
    for name, loc_id in expected.items():
        actual = data.get(name, {}).get("location_id")
        _record(f"{name} -> locationId={loc_id}", actual == loc_id, f"got {actual}")


def test_url_generation():
    """Test that URLs use the modern locationId + category-code format."""
    print("\n[5] URL generation")
    from scraper.utils import generate_search_url, generate_all_category_urls, build_page_url

    # Page 1 with locationId
    url = generate_search_url("bayern", 1, "5510")
    expected = ("https://www.kleinanzeigen.de/s-immobilien/bayern/c195l5510"
                "?posterType=PRIVATE&sortingField=SORTING_DATE")
    _record(f"page 1 with locationId", url == expected, f"got {url}")

    # Page 2 should use ?o=2 appended to the default query
    url = generate_search_url("bayern", 2, "5510")
    expected = ("https://www.kleinanzeigen.de/s-immobilien/bayern/c195l5510"
                "?posterType=PRIVATE&sortingField=SORTING_DATE&o=2")
    _record(f"page 2 uses ?o=", url == expected, f"got {url}")

    # The default query params must always be present (so private-only + sort-by-date)
    _record("default query contains posterType=PRIVATE",
            "posterType=PRIVATE" in url)
    _record("default query contains sortingField=SORTING_DATE",
            "sortingField=SORTING_DATE" in url)

    # Sub-category URLs should use the modern category-code scheme
    urls = generate_all_category_urls("bremen", "1")
    _record("generate_all_category_urls returns URLs",
            len(urls) > 0, f"{len(urls)} URLs")
    # Each URL should have the region slug, a category code, locationId, and the default params
    sample = urls[0] if urls else ""
    _record("URL uses /bremen/c<code>l<locationId> format",
            "/bremen/c" in sample and "l1" in sample,
            f"sample: {sample}")
    _record("URL has posterType=PRIVATE filter",
            all("posterType=PRIVATE" in u for u in urls))
    _record("URL has sortingField=SORTING_DATE",
            all("sortingField=SORTING_DATE" in u for u in urls))

    # Sub-categories should include c208 (Häuser zum Kauf), where the user's
    # example listing lives.
    cat_codes = [u.split("/")[-1].split("l")[0].split("?")[0] for u in urls]
    _record("subcategory list contains c208 (Häuser zum Kauf)",
            "c208" in cat_codes, f"codes: {cat_codes}")

    # build_page_url should produce paginated URL
    base = ("https://www.kleinanzeigen.de/s-immobilien/bremen/c195l1"
            "?posterType=PRIVATE&sortingField=SORTING_DATE")
    p2 = build_page_url(base, 2)
    _record("build_page_url page 2", p2 == base + "&o=2", f"got {p2}")
    _record("build_page_url page 1 == base", build_page_url(base, 1) == base)


def test_min_age_filter():
    """Test that Listing.is_older_than_3_months works."""
    print("\n[6] Age filter")
    from scraper.models import Listing
    from datetime import datetime, timedelta

    fresh = Listing(title="t", url="u", date_parsed=datetime.now() - timedelta(days=10))
    old = Listing(title="t", url="u", date_parsed=datetime.now() - timedelta(days=120))
    boundary = Listing(title="t", url="u", date_parsed=datetime.now() - timedelta(days=91))
    nodate = Listing(title="t", url="u")

    _record("10-day-old listing is not old", not fresh.is_older_than_3_months)
    _record("120-day-old listing IS old", old.is_older_than_3_months)
    _record("91-day-old listing IS old", boundary.is_older_than_3_months)
    _record("listing with no date is not old", not nodate.is_older_than_3_months)


def test_price_threshold():
    """Test the price-floor filter at the configured MIN_PRICE_EUR.

    The rule (effective 2026-07-03, lowered to 100 000 on the same day):
    drop everything whose lowest numeric price is at or below
    ``Settings.MIN_PRICE_EUR`` (currently 100 000), VB or no VB. Drop
    listings with no parsable price. Tests pass a small threshold
    (100) so boundary cases do not need >100 000 fixtures.
    """
    print("\n[7] Price floor")
    from scraper.utils import is_excluded_by_price
    from config.settings import Settings

    # Threshold of 100 makes boundary cases easy to construct. The
    # production default of 100 000 is exercised separately below.
    THR = 100

    cases = [
        # (raw_price, expected_excluded_at_thr=100, description)
        (None,                   True,  "missing price -> exclude"),
        ("",                     True,  "empty string -> exclude"),
        ("VB",                   True,  "VB alone -> exclude"),
        (f"{THR - 1} €",         True,  "below floor -> exclude"),
        (f"{THR} €",             True,  "exactly at floor -> exclude"),
        (f"{THR + 1} €",         False, "one above floor -> keep"),
        (f"{THR * 10} €",        False, "comfortably above floor -> keep"),
        (f"{THR * 10} € VB",     False, "above-floor with VB -> keep"),
        (f"{THR // 2} € VB",     True,  "below-floor with VB -> exclude"),
        # Two-price ranges use the LOWER bound:
        (f"{THR + 5} € VB {THR * 10} €", False, "two prices, lower above floor -> keep"),
        (f"{THR // 2} € VB {THR * 10} €", True,  "two prices, lower at-or-below floor -> exclude"),
    ]
    for raw, expected_excluded, desc in cases:
        got = is_excluded_by_price(raw, min_price_eur=THR)
        _record(desc, got == expected_excluded,
                f"got excluded={got}, expected {expected_excluded}")

    # Production default: must drop the same '1 €' rows that were
    # slipping through before this fix (finding A4 in REVIEW.md).
    _record("production default Settings.MIN_PRICE_EUR is 100000",
            Settings.MIN_PRICE_EUR == 100000, str(Settings.MIN_PRICE_EUR))
    _record("'1 €' is excluded at the default 100 000 floor",
            is_excluded_by_price("1 €"))
    _record("'100.000 €' is excluded at the default 100 000 floor (boundary)",
            is_excluded_by_price("100.000 €"))
    _record("'100.001 €' is kept at the default 100 000 floor",
            not is_excluded_by_price("100.001 €"))
    _record("'200.000 €' is kept at the default 100 000 floor",
            not is_excluded_by_price("200.000 €"))


def test_plz_sort():
    """Test the Excel sort order: activation date ASCENDING (oldest first).

    Effective 2026-07-03 per user request "sort the listings after
    their time since online ... descending (oldest at the top)". The
    sort is purely client-side because Kleinanzeigen has no
    age-descending URL parameter. Empty / unparseable activation date
    sorts to the end so dated rows stay grouped at the top.
    """
    print("\n[8] Excel sort order")
    from datetime import datetime
    from scraper.exporter import _sort_listings, ExcelExporter
    from scraper.models import Listing

    def L(date_str):
        """Build a Listing whose date_parsed matches the date string."""
        if date_str:
            d = datetime.strptime(date_str, "%d.%m.%Y")
            return Listing(title="x", url=f"u-{date_str}",
                           postleitzahl="52072",
                           date_parsed=d,
                           date_posted=date_str)
        return Listing(title="x", url="u-nodate",
                       postleitzahl="52072",
                       date_parsed=None, date_posted="")

    # Shuffled input on purpose. date_parsed equals the parsed date
    # string so date-ascending ordering matches URL ordering 1:1.
    inp = [
        L("28.12.2025"),   # row 0
        L("17.02.2025"),   # row 1
        L("03.02.2026"),   # row 2
        L("01.03.2026"),   # row 3
        L(""),             # row 4: no date -> last
        L("07.01.2025"),   # row 5
    ]
    # Date ASCENDING (oldest first): 07.01.2025, 17.02.2025,
    # 28.12.2025, 03.02.2026, 01.03.2026, then "" (missing) last.
    out = _sort_listings(inp)

    actual_urls = [L.url for L in out]
    expected_urls = ["u-07.01.2025", "u-17.02.2025", "u-28.12.2025",
                     "u-03.02.2026", "u-01.03.2026", "u-nodate"]
    _record("date asc + missing-last ordering",
            actual_urls == expected_urls,
            f"got {actual_urls}")

    # Idempotence: sorting twice == sorting once.
    _record("_sort_listings is idempotent",
            [L.url for L in _sort_listings(out)] == expected_urls)

    # Non-mutation: input list is not mutated.
    inp_copy = list(inp)
    _sort_listings(inp)
    _record("_sort_listings does not mutate input",
            [L.url for L in inp] == [L.url for L in inp_copy])

    # Boundary: single row stays put.
    single = _sort_listings([L("01.01.2025")])
    _record("single-row sort returns the same row",
            len(single) == 1 and single[0].url == "u-01.01.2025")

    # Boundary: empty input.
    _record("empty list -> empty list",
            _sort_listings([]) == [])

    # Row-dict sort key (the global-accumulator path).
    rows = [
        {"Postleitzahl": "52072", "Datum seit online": "01.03.2026"},
        {"Postleitzahl": "52072", "Datum seit online": "28.12.2025"},
        {"Postleitzahl": "52072", "Datum seit online": "17.02.2025"},
        {"Postleitzahl": "52072", "Datum seit online": "07.01.2025"},
        {"Postleitzahl": "52072", "Datum seit online": "03.02.2026"},
        {"Postleitzahl": "52072", "Datum seit online": ""},
    ]
    rows_sorted = sorted(rows, key=ExcelExporter._row_sort_key)
    actual_dates = [r["Datum seit online"] for r in rows_sorted]
    expected_dates = ["07.01.2025", "17.02.2025", "28.12.2025",
                      "03.02.2026", "01.03.2026", ""]
    _record("row-dict sort produces date asc + missing-last ordering",
            actual_dates == expected_dates,
            f"got {actual_dates}")


def test_previous_price_and_columns():
    """Test the previous-price column and the new column layout.

    Effective 2026-07-03 per user request: when the search-card price
    cell shows two prices ("215.000 € VB 265.000 €" or "550 € 700 €"),
    put the higher one in a separate column at the end
    ("Vorheriger Wert (€)"). Single-price cells leave that cell empty.
    """
    print("\n[9] Previous-price column + GLOBAL_COLUMNS layout")
    from scraper.utils import parse_price_pair_eur
    from scraper.models import Listing
    from scraper.exporter import ExcelExporter
    from datetime import datetime
    _exporter = ExcelExporter()  # for the _listing_row instance method

    # ---- Helper: parse_price_pair_eur ----
    pair_cases = [
        ("184.500 €",              184500, None,  False),
        ("1 €",                    1,      None,  False),
        ("215.000 € VB 265.000 €", 215000, 265000, True),
        ("43.900 € VB 59.000 €",   43900,  59000,  True),
        ("550 € 700 €",            550,    700,    False),
        ("150.001 € VB 200.000 €", 150001, 200000, True),
        ("150 € VB",               150,    None,   True),   # single + VB
        ("VB",                     None,   None,   True),
        (None,                     None,   None,   False),
        ("",                       None,   None,   False),
    ]
    for raw, exp_lower, exp_higher, exp_vb in pair_cases:
        lower, higher, vb = parse_price_pair_eur(raw)
        ok = (lower == exp_lower and higher == exp_higher and vb == exp_vb)
        _record(f"parse_price_pair_eur({raw!r}) == ({exp_lower}, {exp_higher}, {exp_vb})",
                ok, f"got ({lower}, {higher}, {vb})")

    # ---- Listing dataclass has the new field ----
    _record("Listing has previous_price_eur field",
            "previous_price_eur" in Listing.__dataclass_fields__)

    # ---- GLOBAL_COLUMNS ends with the new column ----
    _record("GLOBAL_COLUMNS last entry is 'Vorheriger Wert (€)'",
            ExcelExporter.GLOBAL_COLUMNS[-1] == "Vorheriger Wert (€)",
            f"got {ExcelExporter.GLOBAL_COLUMNS}")
    _record("GLOBAL_COLUMNS has 7 entries (was 6)",
            len(ExcelExporter.GLOBAL_COLUMNS) == 7,
            f"got {len(ExcelExporter.GLOBAL_COLUMNS)}")

    # ---- _listing_row writes the new column ----
    L_two = Listing(title="t", url="u",
                    price="215.000 € VB 265.000 €",
                    price_eur=215000.0, previous_price_eur=265000.0,
                    location="city", seller_name="s",
                    date_parsed=datetime(2025, 6, 13),
                    date_posted="13.06.2025",
                    is_older_than_3_months=True)
    row_two = _exporter._listing_row(L_two)
    _record("two-price row has '215,000' in Aktueller Wert",
            row_two["Aktueller Wert (€)"] == "215,000",
            f"got {row_two['Aktueller Wert (€)']!r}")
    _record("two-price row has '265,000' in Vorheriger Wert",
            row_two["Vorheriger Wert (€)"] == "265,000",
            f"got {row_two['Vorheriger Wert (€)']!r}")

    L_single = Listing(title="t", url="u",
                       price="184.500 €",
                       price_eur=184500.0, previous_price_eur=None,
                       location="city", seller_name="s",
                       date_parsed=datetime(2025, 6, 13),
                       date_posted="13.06.2025",
                       is_older_than_3_months=True)
    row_single = _exporter._listing_row(L_single)
    _record("single-price row has '184,500' in Aktueller Wert",
            row_single["Aktueller Wert (€)"] == "184,500",
            f"got {row_single['Aktueller Wert (€)']!r}")
    _record("single-price row has EMPTY Vorheriger Wert (not a repeat)",
            row_single["Vorheriger Wert (€)"] == "",
            f"got {row_single['Vorheriger Wert (€)']!r}")


def test_export_always_writes():
    """Test the always-write behaviour effective 2026-07-03.

    User request: "The next execution of this program should find all
    listings as if it is executed the first time ever. Then it should
    write alle findings in a apropriate excel file." Default is now
    write_all=True (every listing written) and the file is always
    created, even when zero new rows survive dedup. Tested at the
    export-global level using a temp xlsx.
    """
    print("\n[10] Always-write export behaviour")
    import tempfile
    from datetime import datetime, timedelta
    from pathlib import Path
    from openpyxl import load_workbook
    from scraper.exporter import ExcelExporter, export_to_excel
    from scraper.models import Listing, ScrapeResult

    def L(url, age_days):
        """Build a Listing with a date_parsed that yields age_days."""
        if age_days is None:
            return Listing(title="t", url=url,
                           postleitzahl="52072", date_parsed=None,
                           date_posted="", price_eur=215000.0,
                           previous_price_eur=None,
                           location="c", seller_name="s",
                           is_older_than_3_months=False)
        d = datetime(2026, 7, 3) - timedelta(days=age_days)
        return Listing(title="t", url=url,
                       postleitzahl="52072", date_parsed=d,
                       date_posted=d.strftime("%d.%m.%Y"),
                       price_eur=215000.0, previous_price_eur=None,
                       location="c", seller_name="s",
                       is_older_than_3_months=age_days > 90)

    with tempfile.TemporaryDirectory(prefix="always_write_") as td:
        # Inject our tempdir into Settings so export_to_excel creates
        # the test xlsx there, NOT in the live data/output/. The
        # user's intent on 2026-07-03 was that the next real run
        # should "find all listings as if it is executed the first
        # time ever" — we MUST NOT pollute the live xlsx with our
        # synthetic test data, otherwise the next real scrape would
        # skip the real URLs (they'd be "new" but the file would
        # already contain 3 fake rows from this test).
        from config.settings import Settings
        real_output_dir = Settings.OUTPUT_DIR
        real_global_filename = Settings.GLOBAL_FILENAME
        Settings.OUTPUT_DIR = Path(td)
        Settings.GLOBAL_FILENAME = "test.xlsx"

        try:
            # ---- Case 1: fresh run, all old listings, default ----
            listings = [L(f"https://x/u{i}", a)
                        for i, a in enumerate([100, 200, 365])]
            result = ScrapeResult(bundesland="X",
                                  total_listings_found=3,
                                  old_listings_found=3,
                                  listings=listings,
                                  start_time=datetime(2026, 7, 3),
                                  end_time=datetime(2026, 7, 3, 0, 0, 1))
            # Default export_to_excel is allow_all_fallback=True.
            filepath = export_to_excel(result)
            _record("export_to_excel returns a path on a fresh run with all-old listings",
                    bool(filepath), f"got {filepath!r}")
            _record("filepath is under the tempdir (live xlsx NOT polluted)",
                    bool(filepath) and str(filepath).startswith(td),
                    f"got {filepath!r}")

            # ---- Case 2: fresh run, ZERO qualifying old listings ----
            listings_zero_old = [L(f"https://x/v{i}", a)
                                 for i, a in enumerate([10, 20, 30])]
            result_zero = ScrapeResult(bundesland="Y",
                                       total_listings_found=3,
                                       old_listings_found=0,
                                       listings=listings_zero_old,
                                       start_time=datetime(2026, 7, 3),
                                       end_time=datetime(2026, 7, 3, 0, 0, 1))
            # Use a separate tempdir so we test a TRULY fresh file (no
            # existing baseline to dedup against).
            with tempfile.TemporaryDirectory(prefix="always_write_zero_") as td2:
                Settings.OUTPUT_DIR = Path(td2)
                filepath_zero = ExcelExporter().export_global(result_zero)
                _record("export_global returns a path even when zero old listings qualify",
                        bool(filepath_zero), f"got {filepath_zero!r}")
                if filepath_zero:
                    wb = load_workbook(filepath_zero, data_only=True)
                    ws = wb["Global Old Listings"]
                    header = [c.value for c in ws[1]]
                    _record("fresh file has 7-column header",
                            len(header) == 7, f"got {header}")

            # ---- Case 3: export_global default is write_all=True ----
            _record("ExcelExporter.export_global default write_all is True",
                    ExcelExporter.export_global.__defaults__[0] is True,
                    f"got default={ExcelExporter.export_global.__defaults__[0]!r}")

            # ---- Case 4: export_to_excel default is allow_all_fallback=True ----
            _record("export_to_excel default allow_all_fallback is True",
                    export_to_excel.__defaults__[0] is True,
                    f"got default={export_to_excel.__defaults__[0]!r}")

            # ---- Case 5: dedup still works after the default flip ----
            # Re-run with same URLs: all should be skipped, file still
            # written (with header + Summary, no new rows).
            Settings.OUTPUT_DIR = Path(td)
            filepath_dup = export_to_excel(result)
            if filepath_dup:
                wb = load_workbook(filepath_dup, data_only=True)
                ws = wb["Global Old Listings"]
                rows = list(ws.iter_rows(min_row=2, values_only=True))
                # Strip control rows.
                hdr_local = [c.value for c in ws[1]]
                date_col = hdr_local.index("Datum seit online")
                plz_col = hdr_local.index("Postleitzahl")
                seller_col = hdr_local.index("Name des Verkäufers")
                if (rows and rows[-1][date_col] not in (None, "")
                        and "-" in str(rows[-1][date_col])
                        and len(str(rows[-1][date_col])) == 10
                        and all(v in (None, "") for v in (rows[-1][plz_col], rows[-1][seller_col]))):
                    rows.pop()
                if rows and all(v is None or v == "" for v in rows[-1]):
                    rows.pop()
                _record("dedup correctly skips re-runs (file has 3 rows, not 6)",
                        len(rows) == 3, f"got {len(rows)} rows")
                _record("second run duplicate counter == 3",
                        result.duplicate_global_rows == 3,
                        f"got {result.duplicate_global_rows}")
        finally:
            # Restore Settings so we don't leak state into any other
            # test or any later code.
            Settings.OUTPUT_DIR = real_output_dir
            Settings.GLOBAL_FILENAME = real_global_filename


def main():
    print("=" * 60)
    print("Kleinanzeigen Scraper - Tests")
    print("=" * 60)

    tests = [
        test_imports,
        test_date_parsing,
        test_settings,
        test_bundesland_mapping,
        test_url_generation,
        test_min_age_filter,
        test_price_threshold,
        test_plz_sort,
        test_previous_price_and_columns,
        test_export_always_writes,
    ]
    for t in tests:
        try:
            t()
        except Exception as e:
            _record(t.__name__, False, str(e))
            traceback.print_exc()

    print("\n" + "=" * 60)
    if _failures:
        print(f"❌ {len(_failures)} check(s) FAILED:")
        for name in _failures:
            print(f"   - {name}")
        print("=" * 60)
        return 1
    else:
        print("✅ All checks passed")
        print("=" * 60)
        return 0


if __name__ == "__main__":
    sys.exit(main())